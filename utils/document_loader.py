import os
import json
import logging
import csv
import re
import hashlib
from pathlib import Path
from typing import List, Dict, Any, Optional
from PIL import Image

import config
from utils.ocr import OCRManager

logger = logging.getLogger(__name__)


class DocumentLoader:
    @staticmethod
    def compute_file_hash(file_path: str) -> str:
        """Calculate SHA-256 hash of a file for cache invalidation and deduplication."""
        hasher = hashlib.sha256()
        try:
            with open(file_path, "rb") as f:
                while chunk := f.read(65536):
                    hasher.update(chunk)
            return hasher.hexdigest()
        except Exception as e:
            logger.warning(f"Could not compute hash for {file_path}: {e}")
            return ""

    @staticmethod
    def load_pdf_pages(file_path: str) -> List[Dict[str, Any]]:
        """Extract text page by page from PDF.
        
        Uses PyMuPDF (fitz) for fast native extraction.
        Detects pages with insufficient native text and selectively applies OCR only to those pages.
        Supports mixed PDFs (digital + scanned) while strictly preserving page numbers and reading order.
        Handles empty, corrupted, and password-protected files gracefully.
        """
        pages = []
        path_obj = Path(file_path)
        file_name = path_obj.name

        if not path_obj.exists() or path_obj.stat().st_size == 0:
            logger.warning(f"PDF file is empty or missing: {file_path}")
            return [{"page_number": 1, "text": "", "is_ocr": False, "file_name": file_name}]

        # 1. Primary engine: PyMuPDF (fitz)
        try:
            import fitz
            doc = fitz.open(file_path)

            if doc.is_encrypted:
                logger.warning(f"PDF is encrypted: {file_path}")
                # Try empty password authenticate
                if not doc.authenticate(""):
                    return [{
                        "page_number": 1,
                        "text": f"Error: '{file_name}' is password protected and cannot be read without credentials.",
                        "is_ocr": False,
                        "file_name": file_name
                    }]

            total_pages = len(doc)
            if total_pages == 0:
                return [{"page_number": 1, "text": "", "is_ocr": False, "file_name": file_name}]

            for page_num in range(total_pages):
                page_idx = page_num + 1
                try:
                    page = doc[page_num]
                    native_text = page.get_text("text") or ""
                    clean_native = native_text.strip()
                    alpha_count = len(re.findall(r'[a-zA-Z0-9]', clean_native))

                    # If page has sufficient usable digital text, use it directly
                    if alpha_count >= 20:
                        pages.append({
                            "page_number": page_idx,
                            "text": clean_native,
                            "is_ocr": False,
                            "file_name": file_name
                        })
                    else:
                        # Insufficient digital text -> Apply OCR specifically to this page
                        logger.info(f"Page {page_idx} of {file_name} has insufficient native text ({alpha_count} chars). Applying OCR...")
                        pix = page.get_pixmap(dpi=150)
                        pil_img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                        ocr_text = OCRManager.extract_text_from_pil(pil_img, preprocess=True, return_warning=False)

                        final_page_text = ocr_text.strip() if ocr_text.strip() else clean_native
                        pages.append({
                            "page_number": page_idx,
                            "text": final_page_text,
                            "is_ocr": bool(ocr_text.strip()),
                            "file_name": file_name
                        })
                except Exception as page_err:
                    logger.error(f"Error extracting page {page_idx} of {file_name}: {page_err}")
                    pages.append({
                        "page_number": page_idx,
                        "text": f"[Error reading page {page_idx}]",
                        "is_ocr": False,
                        "file_name": file_name
                    })

            doc.close()
            return pages

        except Exception as fitz_err:
            logger.warning(f"PyMuPDF failed on {file_path}: {fitz_err}. Falling back to pdfplumber / pypdf.")

        # 2. Secondary fallback: pdfplumber & pypdf
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for page_idx, page in enumerate(pdf.pages, start=1):
                    native_text = page.extract_text() or ""
                    clean_native = native_text.strip()
                    alpha_count = len(re.findall(r'[a-zA-Z0-9]', clean_native))

                    if alpha_count >= 20:
                        pages.append({
                            "page_number": page_idx,
                            "text": clean_native,
                            "is_ocr": False,
                            "file_name": file_name
                        })
                    else:
                        # Render to image for OCR
                        try:
                            im = page.to_image(resolution=150)
                            pil_img = im.original
                            ocr_text = OCRManager.extract_text_from_pil(pil_img, preprocess=True)
                            final_text = ocr_text.strip() if ocr_text.strip() else clean_native
                            pages.append({
                                "page_number": page_idx,
                                "text": final_text,
                                "is_ocr": bool(ocr_text.strip()),
                                "file_name": file_name
                            })
                        except Exception as ocr_err:
                            logger.error(f"Fallback OCR failed on page {page_idx}: {ocr_err}")
                            pages.append({
                                "page_number": page_idx,
                                "text": clean_native,
                                "is_ocr": False,
                                "file_name": file_name
                            })
            if pages:
                return pages
        except Exception as plum_err:
            logger.error(f"pdfplumber failed for {file_path}: {plum_err}")

        # 3. Tertiary fallback: pypdf pure digital
        try:
            import pypdf
            with open(file_path, "rb") as f:
                reader = pypdf.PdfReader(f)
                for page_idx, page in enumerate(reader.pages, start=1):
                    p_text = page.extract_text() or ""
                    pages.append({
                        "page_number": page_idx,
                        "text": p_text.strip(),
                        "is_ocr": False,
                        "file_name": file_name
                    })
            return pages
        except Exception as pypdf_err:
            logger.error(f"All PDF extractors failed for {file_path}: {pypdf_err}")
            return [{
                "page_number": 1,
                "text": f"Error: Corrupted or unreadable PDF file '{file_name}'.",
                "is_ocr": False,
                "file_name": file_name
            }]

    @staticmethod
    def load_pdf(file_path: str) -> str:
        """Extract full PDF text preserving page boundaries and OCR fallback."""
        pages = DocumentLoader.load_pdf_pages(file_path)
        formatted_chunks = []
        for p in pages:
            ocr_tag = " (OCR Scanned)" if p.get("is_ocr") else ""
            if p["text"]:
                formatted_chunks.append(f"--- Page {p['page_number']}{ocr_tag} ---\n{p['text']}")
        return "\n\n".join(formatted_chunks).strip()

    @staticmethod
    def load_image(file_path: str) -> str:
        """Extract text from an image file using OCR."""
        file_name = Path(file_path).name
        ocr_text = OCRManager.extract_text_from_image(file_path, preprocess=True, return_warning=False)
        if ocr_text and not ocr_text.startswith("Error") and not ocr_text.startswith("⚠️"):
            return f"--- OCR Extracted Text from Image: {file_name} ---\n{ocr_text}"
        return ""

    @staticmethod
    def load_docx(file_path: str) -> str:
        """Extract text from Word Document preserving paragraphs and tables."""
        import docx
        text_parts = []
        try:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                p_text = para.text.strip()
                if p_text:
                    text_parts.append(p_text)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        text_parts.append(" | ".join(row_cells))
        except Exception as e:
            logger.error(f"python-docx extraction failed for {file_path}: {e}")
        return "\n\n".join(text_parts).strip()

    @staticmethod
    def load_txt(file_path: str) -> str:
        """Extract text from TXT with robust multi-encoding detection."""
        encodings = ['utf-8', 'latin-1', 'cp1252', 'utf-16', 'utf-8-sig']
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    return f.read().strip()
            except (UnicodeDecodeError, LookupError):
                continue
        logger.error(f"Failed to decode text file: {file_path}")
        return ""

    @staticmethod
    def load_md(file_path: str) -> str:
        """Extract text from Markdown file."""
        return DocumentLoader.load_txt(file_path)

    @staticmethod
    def load_csv(file_path: str) -> str:
        """Extract tabular text from CSV file with column metadata."""
        lines = []
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                reader = csv.reader(f)
                headers = next(reader, None)
                if headers:
                    lines.append(f"Table Columns: {', '.join(headers)}")
                for i, row in enumerate(reader, 1):
                    if row:
                        row_str = " | ".join([cell.strip() for cell in row])
                        lines.append(f"Row {i}: {row_str}")
        except Exception as e:
            logger.error(f"CSV extraction failed for {file_path}: {e}")
        return "\n".join(lines)

    @staticmethod
    def load_json(file_path: str) -> str:
        """Extract text from JSON file."""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                data = json.load(f)
                return json.dumps(data, indent=2)
        except Exception as e:
            logger.error(f"JSON extraction failed for {file_path}: {e}")
            return DocumentLoader.load_txt(file_path)

    @classmethod
    def load_document_pages(cls, file_path: str) -> List[Dict[str, Any]]:
        """Return structured page-by-page data for any supported document or image."""
        path = Path(file_path)
        if not path.exists():
            return []

        ext = path.suffix.lower()
        file_name = path.name

        if ext == ".pdf":
            return cls.load_pdf_pages(file_path)
        elif ext in config.SUPPORTED_IMAGE_EXTENSIONS:
            text = cls.load_image(file_path)
            return [{
                "page_number": 1,
                "text": text,
                "is_ocr": True,
                "file_name": file_name
            }]
        else:
            # Single-page / non-paginated documents (docx, txt, md, csv, json)
            text = cls.load_document(file_path)
            return [{
                "page_number": 1,
                "text": text,
                "is_ocr": False,
                "file_name": file_name
            }]

    @classmethod
    def load_document(cls, file_path: str) -> str:
        """Central router for document & image text extraction based on file extension."""
        path = Path(file_path)
        if not path.exists():
            logger.error(f"File does not exist: {file_path}")
            return ""

        ext = path.suffix.lower()
        if ext == ".pdf":
            return cls.load_pdf(file_path)
        elif ext == ".docx":
            return cls.load_docx(file_path)
        elif ext == ".txt":
            return cls.load_txt(file_path)
        elif ext == ".md":
            return cls.load_md(file_path)
        elif ext == ".csv":
            return cls.load_csv(file_path)
        elif ext == ".json":
            return cls.load_json(file_path)
        elif ext in config.SUPPORTED_IMAGE_EXTENSIONS:
            return cls.load_image(file_path)
        else:
            logger.warning(f"Unsupported file extension: {ext}")
            return ""
